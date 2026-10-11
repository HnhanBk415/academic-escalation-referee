import io
from copy import deepcopy
from pathlib import Path

import pytest
import yaml
from docx import Document

from app.core.errors import AppError
from app.services.format_policies import validate_format_policy_document


def _docx_with_heading(heading: str) -> bytes:
    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph("Nội dung báo cáo mẫu.")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _docx_with_cover_and_headings(headings: list[str], include_cover: bool = True) -> bytes:
    doc = Document()
    if include_cover:
        doc.add_paragraph("ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH\nTRƯỜNG ĐẠI HỌC BÁCH KHOA\nKHOA KHOA HỌC VÀ KỸ THUẬT MÁY TÍNH")
        doc.add_paragraph("BÁO CÁO ĐỒ ÁN MÔN HỌC")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Giảng viên hướng dẫn: TS. Nguyễn Văn B"
        table.rows[0].cells[1].text = "Học kỳ: HK242"
        table.rows[1].cells[0].text = "Sinh viên thực hiện: Lê Văn C"
        table.rows[1].cells[1].text = "MSSV: 2011003"
        doc.add_paragraph("TP. HỒ CHÍ MINH, NĂM 2026")
        doc.add_page_break()

    for h in headings:
        doc.add_heading(h, level=1 if "Chương" in h else 2)
        doc.add_paragraph(f"Nội dung chi tiết của phần {h}.")

    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _demo_policy() -> tuple[dict, Path]:
    path = Path(__file__).resolve().parents[2] / "data" / "policies" / "CO3001" / "261.yaml"
    return yaml.safe_load(path.read_text(encoding="utf-8")), path


def test_policy_source_contract_accepts_demo_policy() -> None:
    policy, path = _demo_policy()
    assert validate_format_policy_document(policy, path)["policy_id"] == "format-policy/CO3001/261"


def test_policy_source_contract_rejects_duplicate_rule_ids() -> None:
    policy, path = _demo_policy()
    invalid_policy = deepcopy(policy)
    rules = invalid_policy["report_profiles"][1]["rules"]
    duplicate = deepcopy(rules[0])
    duplicate["title"] = "Bản sao rule không hợp lệ"
    rules.append(duplicate)

    with pytest.raises(AppError, match="bị lặp") as error:
        validate_format_policy_document(invalid_policy, path)

    assert error.value.code == "FORMAT_POLICY_INVALID"


def test_policy_source_contract_rejects_manual_rule_that_cannot_reach_review() -> None:
    policy, path = _demo_policy()
    invalid_policy = deepcopy(policy)
    rule = invalid_policy["report_profiles"][1]["rules"][-1]
    rule["handling"] = "BLOCKING"
    rule["when_unable_to_check"] = "UNABLE_TO_CHECK"

    with pytest.raises(AppError, match="manual"):
        validate_format_policy_document(invalid_policy, path)


def test_policy_source_contract_rejects_tampered_source_checksum() -> None:
    policy, path = _demo_policy()
    invalid_policy = deepcopy(policy)
    invalid_policy["source_documents"][0]["sha256"] = "0" * 64

    with pytest.raises(AppError, match="checksum"):
        validate_format_policy_document(invalid_policy, path)


@pytest.mark.asyncio
async def test_format_check_requires_approved_criteria(client):
    response = await client.post(
        "/api/v1/format-checks/checks",
        data={"course_id": "CO3001", "actor_id": "demo-student"},
        files={"file": ("report.docx", _docx_with_heading("Mở đầu"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    assert response.json()["overall_status"] == "UNABLE_TO_CHECK"
    assert response.json()["results"] == []


@pytest.mark.asyncio
async def test_template_criteria_is_approved_and_checked_without_ai(client):
    created = await client.post(
        "/api/v1/format-checks/criteria/drafts",
        json={
            "course_id": "CO3001",
            "lecturer_id": "lecturer-01",
            "title": "Mẫu báo cáo",
            "criteria": [
                {
                    "name": "Khung đề mục",
                    "description": "Bài phải có đề mục Mở đầu.",
                    "rule_type": "TEMPLATE",
                    "expected_value": {"headings": ["Mở đầu"], "allow_extra": True, "strict_names": True},
                }
            ],
        },
    )
    assert created.status_code == 201, created.text
    criteria_id = created.json()["id"]
    assert created.json()["criteria"][0]["automatable"] is True

    approved = await client.post(f"/api/v1/format-checks/criteria/{criteria_id}/approve?lecturer_id=lecturer-01")
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "ACTIVE"

    response = await client.post(
        "/api/v1/format-checks/checks",
        data={"course_id": "CO3001", "actor_id": "demo-student"},
        files={"file": ("report.docx", _docx_with_heading("Mở đầu"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200, response.text
    assert response.json()["overall_status"] == "PASS"
    assert response.json()["criteria_version"] == 1
    assert response.json()["results"][0]["status"] == "PASS"


@pytest.mark.asyncio
async def test_policy_is_selected_by_report_type_and_keeps_versions_separate(client):
    async def create_active(report_type: str, title: str) -> dict:
        draft = await client.post(
            "/api/v1/format-checks/criteria/drafts",
            json={
                "course_id": "CO3001",
                "lecturer_id": "lecturer-01",
                "report_type": report_type,
                "title": title,
                "criteria": [{
                    "name": "Mở đầu",
                    "description": "Bài có đề mục Mở đầu.",
                    "rule_type": "TEMPLATE",
                    "expected_value": {"headings": ["Mở đầu"], "allow_extra": True},
                }],
            },
        )
        assert draft.status_code == 201, draft.text
        approved = await client.post(
            f"/api/v1/format-checks/criteria/{draft.json()['id']}/approve?lecturer_id=lecturer-01"
        )
        assert approved.status_code == 200, approved.text
        return approved.json()

    # The YAML source is imported by seed with its own policy metadata.
    seeded_final = await client.get(
        "/api/v1/format-checks/criteria?course_id=CO3001&report_type=final_project_report"
    )
    assert seeded_final.status_code == 200
    assert any(
        item["status"] == "ACTIVE"
        and item["policy_id"] == "format-policy/CO3001/261"
        and item["policy_version"] == "1.0.0"
        for item in seeded_final.json()
    )

    progress_v1 = await create_active("progress_report", "Báo cáo tiến độ")
    final_v1 = await create_active("final_project_report", "Báo cáo cuối kỳ")

    progress_check = await client.post(
        "/api/v1/format-checks/checks",
        data={"course_id": "CO3001", "actor_id": "demo-student", "report_type": "progress_report"},
        files={"file": ("progress.docx", _docx_with_heading("Mở đầu"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert progress_check.status_code == 200, progress_check.text
    assert progress_check.json()["criteria_set_id"] == progress_v1["id"]
    assert progress_check.json()["report_type"] == "progress_report"
    assert progress_check.json()["policy_id"] == "format-policy/CO3001/261"
    assert progress_check.json()["policy_version"] == progress_v1["policy_version"]
    assert progress_check.json()["check_run_id"]

    # Activating a new progress policy must not supersede the final-report policy.
    progress_v2 = await create_active("progress_report", "Báo cáo tiến độ v2")
    criteria = await client.get("/api/v1/format-checks/criteria?course_id=CO3001")
    assert criteria.status_code == 200
    active_by_type = {
        item["report_type"]: item["id"]
        for item in criteria.json()
        if item["status"] == "ACTIVE"
    }
    assert active_by_type == {
        "progress_report": progress_v2["id"],
        "final_project_report": final_v1["id"],
    }


@pytest.mark.asyncio
async def test_rule_handling_controls_precheck_outcome(client):
    async def activate(report_type: str, criterion: dict) -> None:
        draft = await client.post(
            "/api/v1/format-checks/criteria/drafts",
            json={
                "course_id": "CO3001",
                "lecturer_id": "lecturer-01",
                "report_type": report_type,
                "title": report_type,
                "criteria": [criterion],
            },
        )
        assert draft.status_code == 201, draft.text
        approved = await client.post(
            f"/api/v1/format-checks/criteria/{draft.json()['id']}/approve?lecturer_id=lecturer-01"
        )
        assert approved.status_code == 200, approved.text

    await activate(
        "advisory_report",
        {
            "name": "Khuyến nghị có mục Kết luận",
            "description": "Nên có mục Kết luận.",
            "rule_type": "TEMPLATE",
            "handling": "ADVISORY",
            "expected_value": {
                "kind": "template_outline",
                "outline_headings": ["Kết luận"],
                "allow_extra": True,
            },
        },
    )
    advisory = await client.post(
        "/api/v1/format-checks/checks",
        data={"course_id": "CO3001", "actor_id": "demo-student", "report_type": "advisory_report"},
        files={"file": ("report.docx", _docx_with_heading("Mở đầu"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert advisory.status_code == 200, advisory.text
    assert advisory.json()["overall_status"] == "PASS"
    assert advisory.json()["results"][0]["status"] == "FAIL"
    assert advisory.json()["results"][0]["handling"] == "ADVISORY"
    assert advisory.json()["review_case_id"] is None

    await activate(
        "manual_report",
        {
            "name": "Chuẩn trích dẫn IEEE",
            "description": "Giảng viên xác nhận cách trích dẫn.",
            "rule_type": "PRESENCE",
            "handling": "MANUAL_REVIEW",
            "review_only": True,
            "expected_value": {"kind": "reference_style", "value": "IEEE"},
        },
    )
    manual = await client.post(
        "/api/v1/format-checks/checks",
        data={"course_id": "CO3001", "actor_id": "demo-student", "report_type": "manual_report"},
        files={"file": ("report.docx", _docx_with_heading("Mở đầu"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert manual.status_code == 200, manual.text
    assert manual.json()["overall_status"] == "NEEDS_REVIEW"
    assert manual.json()["results"][0]["status"] == "MANUAL_REVIEW"
    assert manual.json()["review_case_id"]

    await activate(
        "unavailable_report",
        {
            "name": "Xác nhận thủ công bắt buộc",
            "description": "Không có dữ liệu tự động để đo.",
            "rule_type": "PRESENCE",
            "handling": "BLOCKING",
            "review_only": True,
            "when_unable_to_check": "UNABLE_TO_CHECK",
            "expected_value": {"kind": "reference_style", "value": "IEEE"},
        },
    )
    unavailable = await client.post(
        "/api/v1/format-checks/checks",
        data={"course_id": "CO3001", "actor_id": "demo-student", "report_type": "unavailable_report"},
        files={"file": ("report.docx", _docx_with_heading("Mở đầu"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert unavailable.status_code == 200, unavailable.text
    assert unavailable.json()["overall_status"] == "UNABLE_TO_CHECK"
    assert unavailable.json()["review_case_id"] is None


@pytest.mark.asyncio
async def test_ambiguous_template_creates_limited_review_case(client):
    created = await client.post(
        "/api/v1/format-checks/criteria/drafts",
        json={
            "course_id": "CO3001",
            "lecturer_id": "lecturer-01",
            "title": "Mẫu báo cáo",
            "criteria": [{
                "name": "Khung đề mục", "description": "Có mục Tổng quan.", "rule_type": "TEMPLATE",
                "expected_value": {
                    "kind": "template_outline",
                    "outline_headings": ["Tổng quan học phần"],
                    "allow_extra": True,
                    "strict_names": True,
                },
            }],
        },
    )
    criteria_id = created.json()["id"]
    await client.post(f"/api/v1/format-checks/criteria/{criteria_id}/approve?lecturer_id=lecturer-01")
    checked = await client.post(
        "/api/v1/format-checks/checks",
        data={"course_id": "CO3001", "actor_id": "demo-student"},
        files={"file": ("report.docx", _docx_with_heading("Tổng quan đề tài"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert checked.status_code == 200, checked.text
    assert checked.json()["overall_status"] == "NEEDS_REVIEW"
    review_id = checked.json()["review_case_id"]
    assert review_id

    reviews = await client.get("/api/v1/format-checks/reviews?course_id=CO3001&lecturer_id=lecturer-01")
    assert reviews.status_code == 200
    assert reviews.json()[0]["id"] == review_id
    assert "headings" in reviews.json()[0]["extracted_snapshot"]

    decided = await client.post(
        f"/api/v1/format-checks/reviews/{review_id}/decision",
        json={"reviewer_id": "lecturer-01", "decision": "ACCEPTED", "reason": "Tên đề mục tương đương."},
    )
    assert decided.status_code == 200
    assert decided.json()["status"] == "ACCEPTED"


@pytest.mark.asyncio
async def test_academic_30_checklist_draft_and_evaluation(client):
    # 1. Trích để tạo checklist từ quy định: cái nào có thì nạp, không có bỏ trống
    draft_res = await client.post(
        "/api/v1/format-checks/criteria/drafts/from-documents",
        json={
            "course_id": "CO3001",
            "lecturer_id": "lecturer-01",
            "title": "Checklist Quy chế Báo cáo Đồ án",
            "document_ids": ["co3001-format-guidelines-v1"],
            "include_template_structure": True,
        },
    )
    assert draft_res.status_code == 201, draft_res.text
    draft_data = draft_res.json()
    criteria = draft_data["criteria"]
    assert len(criteria) >= 4

    # Verify populated items from regulation
    criteria_names = [c["name"] for c in criteria]
    assert any("Khổ giấy" in name for name in criteria_names)
    assert any("Phông chữ" in name for name in criteria_names)

    # Approve criteria
    crit_set_id = draft_data["id"]
    await client.post(f"/api/v1/format-checks/criteria/{crit_set_id}/approve?lecturer_id=lecturer-01")

    # 2 & 3. Trích facts từ báo cáo và giao cả 2 cho AI / Rule engine đánh giá
    checked = await client.post(
        "/api/v1/format-checks/checks",
        data={"course_id": "CO3001", "actor_id": "demo-student"},
        files={"file": ("report_sv.docx", _docx_with_heading("1. Mô tả đồ án"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert checked.status_code == 200, checked.text
    res_data = checked.json()

    # Bảng tổng hợp kiểm tra động (Checklist summary tương ứng đúng số lượng tiêu chí)
    summary = res_data.get("checklist_summary")
    assert summary is not None
    assert len(summary) == len(criteria)

    # Tiêu chí có quy định được đánh giá (PASS / FAIL / AMBIGUOUS)
    page_setup_item = next(item for item in summary if "Khổ giấy" in item["category"] or "Căn lề" in item["category"])
    assert page_setup_item["status"] in {"PASS", "FAIL"}


@pytest.mark.asyncio
async def test_cover_page_detection_and_evaluation(client):
    # 1. Create criteria set with Cover Page
    draft_res = await client.post(
        "/api/v1/format-checks/criteria/drafts/from-documents",
        json={
            "course_id": "CO3001",
            "lecturer_id": "lecturer-01",
            "title": "Quy chế Báo cáo có Trang bìa",
            "document_ids": ["co3001-format-guidelines-v1"],
            "include_template_structure": True,
        },
    )
    assert draft_res.status_code == 201
    crit_set_id = draft_res.json()["id"]
    await client.post(f"/api/v1/format-checks/criteria/{crit_set_id}/approve?lecturer_id=lecturer-01")

    # 2. Submission WITH full cover page
    file_with_cover = _docx_with_cover_and_headings(["Chương 1: Khái niệm", "1. Mô tả đồ án"], include_cover=True)
    checked_pass = await client.post(
        "/api/v1/format-checks/checks",
        data={"course_id": "CO3001", "actor_id": "demo-student"},
        files={"file": ("report_with_cover.docx", file_with_cover, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert checked_pass.status_code == 200
    summary_pass = checked_pass.json()["checklist_summary"]
    cover_item = next(item for item in summary_pass if "bìa" in item["category"].lower() or item.get("index") == 5)
    assert cover_item["status"] == "PASS", cover_item

    # 3. Submission WITHOUT cover page
    file_no_cover = _docx_with_cover_and_headings(["Chương 1: Khái niệm"], include_cover=False)
    checked_fail = await client.post(
        "/api/v1/format-checks/checks",
        data={"course_id": "CO3001", "actor_id": "demo-student"},
        files={"file": ("report_no_cover.docx", file_no_cover, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert checked_fail.status_code == 200
    summary_fail = checked_fail.json()["checklist_summary"]
    cover_item_fail = next(item for item in summary_fail if "bìa" in item["category"].lower() or item.get("index") == 5)
    assert cover_item_fail["status"] in ("FAIL", "MISSING_STUDENT_INPUT")
    assert "trang bìa" in cover_item_fail["error_summary"].lower() or "bổ sung" in cover_item_fail["suggestion"].lower() or "chưa đạt" in cover_item_fail["error_summary"].lower()



@pytest.mark.asyncio
async def test_granular_template_and_semantic_equivalence_toggle(client):
    # A. Draft with natural outline criteria
    draft_semantic = await client.post(
        "/api/v1/format-checks/criteria/drafts/from-documents",
        json={
            "course_id": "CO3001",
            "lecturer_id": "lecturer-01",
            "title": "Checklist Khung cấu trúc đề cương",
            "document_ids": ["co3001-format-guidelines-v1"],
            "include_template_structure": True,
            "allow_semantic_equivalent": True,
        },
    )
    assert draft_semantic.status_code == 201
    crit_semantic = draft_semantic.json()["criteria"]

    # Verify natural text template criteria exist
    names = [c["name"] for c in crit_semantic]
    assert any("Khung cấu trúc" in n for n in names)

    await client.post(f"/api/v1/format-checks/criteria/{draft_semantic.json()['id']}/approve?lecturer_id=lecturer-01")

    # Submission using semantically equivalent heading: "Chương 1: Mở đầu" instead of "Chương 1: Khái niệm"
    sem_doc = _docx_with_cover_and_headings([
        "Chương 1: Mở đầu",
        "1.1 Giới thiệu đề tài",
        "Chương 2: Hiện thực",
        "Chương 3: Đánh giá và Kết luận",
    ])
    res_sem = await client.post(
        "/api/v1/format-checks/checks",
        data={"course_id": "CO3001", "actor_id": "demo-student"},
        files={"file": ("sem_report.docx", sem_doc, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res_sem.status_code == 200
    template_res = next(r for r in res_sem.json()["results"] if "Khung cấu trúc" in r["name"])
    assert template_res["status"] == "PASS"

    # B. Missing required chapters
    incomplete_doc = _docx_with_cover_and_headings([
        "Chương 1: Khái niệm",
    ])
    res_incomplete = await client.post(
        "/api/v1/format-checks/checks",
        data={"course_id": "CO3001", "actor_id": "demo-student"},
        files={"file": ("incomplete.docx", incomplete_doc, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res_incomplete.status_code == 200
    template_inc_res = next(r for r in res_incomplete.json()["results"] if "Khung cấu trúc" in r["name"])
    assert template_inc_res["status"] in ("FAIL", "MISSING_STUDENT_INPUT", "AMBIGUOUS")
@pytest.mark.asyncio
async def test_zero_defaults_strict_extraction():
    from app.ai.fake import FakeProvider

    criteria = await FakeProvider().extract_regulation_criteria(
        regulation_text="Quy chế định dạng...",
        allow_semantic_equivalent=True,
    )
    assert len(criteria) >= 9

    # Verify criteria are dynamic and have exact requirement text extracted from regulation
    for c in criteria:
        assert c.name
        assert c.rule_type
        assert c.description
