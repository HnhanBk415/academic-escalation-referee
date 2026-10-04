import io
import re
from collections import Counter
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.enums import (
    ActorRole,
    FormatCheckOverallStatus,
    FormatCriteriaSetStatus,
    FormatCriterionResultStatus,
    FormatCriterionRuleType,
    FormatReviewStatus,
)
from app.core.errors import AppError
from app.models import (
    Actor,
    Course,
    Document,
    DocumentChunk,
    FormatCheckMetric,
    FormatCriteriaSet,
    FormatCriterion,
    FormatReviewCase,
)
from app.schemas.format_checks import (
    FormatCheckResponse,
    FormatCriteriaDraftCreate,
    FormatCriteriaSetResponse,
    FormatCriterionInput,
    FormatCriterionResponse,
    FormatCriterionResult,
    FormatDraftFromDocumentsCreate,
    FormatMetricResponse,
    FormatReviewCaseResponse,
    FormatReviewDecisionCreate,
)
from app.services.audit import add_audit_event
from app.services.ids import new_id


SUPPORTED_MEASUREMENTS = {
    "page_count",
    "margins_cm",
    "font_name",
    "font_size_pt",
    "line_spacing",
}


def _fold(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().casefold())


def _criterion_is_automatable(item: FormatCriterionInput) -> bool:
    if item.review_only:
        return False
    if item.rule_type == FormatCriterionRuleType.MEASUREMENT:
        return str(item.expected_value.get("property", "")) in SUPPORTED_MEASUREMENTS
    if item.rule_type == FormatCriterionRuleType.PRESENCE:
        return str(item.expected_value.get("kind", "")) in {"heading", "table_columns"}
    if item.rule_type == FormatCriterionRuleType.TEMPLATE:
        return bool(item.expected_value.get("headings"))
    return False


async def _validate_course_lecturer(session: AsyncSession, course_id: str, lecturer_id: str) -> None:
    course = await session.get(Course, course_id)
    lecturer = await session.get(Actor, lecturer_id)
    # Courses currently have one configured responsible lecturer in the demo.
    if course is None:
        raise AppError("COURSE_NOT_FOUND", "Không tìm thấy học phần.", status_code=404)
    if lecturer is None or lecturer.role != ActorRole.LECTURER or lecturer.id != "lecturer-01":
        raise AppError(
            "FORMAT_CRITERIA_NOT_AUTHORIZED",
            "Chỉ giảng viên phụ trách học phần mới được quản lý bộ tiêu chí.",
            status_code=403,
        )


def _criterion_response(item: FormatCriterion) -> FormatCriterionResponse:
    return FormatCriterionResponse(
        id=item.id,
        position=item.position,
        name=item.name,
        description=item.description,
        rule_type=item.rule_type,
        expected_value=item.expected_value,
        source_citation=item.source_citation,
        source_document_id=item.source_document_id,
        review_only=item.review_only,
        automatable=item.automatable,
    )


async def _criteria_response(session: AsyncSession, item: FormatCriteriaSet) -> FormatCriteriaSetResponse:
    criteria = (
        await session.scalars(
            select(FormatCriterion)
            .where(FormatCriterion.criteria_set_id == item.id)
            .order_by(FormatCriterion.position)
        )
    ).all()
    return FormatCriteriaSetResponse(
        id=item.id,
        course_id=item.course_id,
        version=item.version,
        title=item.title,
        status=item.status,
        source_document_ids=item.source_document_ids or [],
        created_by=item.created_by,
        approved_by=item.approved_by,
        approved_at=item.approved_at,
        created_at=item.created_at,
        criteria=[_criterion_response(criterion) for criterion in criteria],
    )


async def create_draft(
    session: AsyncSession, payload: FormatCriteriaDraftCreate, request_id: str | None
) -> FormatCriteriaSetResponse:
    await _validate_course_lecturer(session, payload.course_id, payload.lecturer_id)
    next_version = int(
        await session.scalar(
            select(func.coalesce(func.max(FormatCriteriaSet.version), 0)).where(
                FormatCriteriaSet.course_id == payload.course_id
            )
        )
        or 0
    ) + 1
    criteria_set = FormatCriteriaSet(
        id=new_id("format_set"),
        course_id=payload.course_id,
        version=next_version,
        title=payload.title,
        status=FormatCriteriaSetStatus.DRAFT,
        source_document_ids=payload.source_document_ids,
        created_by=payload.lecturer_id,
    )
    session.add(criteria_set)
    for position, criterion in enumerate(payload.criteria, start=1):
        session.add(
            FormatCriterion(
                id=new_id("format_criterion"),
                criteria_set_id=criteria_set.id,
                position=position,
                name=criterion.name,
                description=criterion.description,
                rule_type=criterion.rule_type,
                expected_value=criterion.expected_value,
                source_citation=criterion.source_citation,
                source_document_id=criterion.source_document_id,
                automatable=_criterion_is_automatable(criterion),
                review_only=criterion.review_only,
            )
        )
    add_audit_event(
        session,
        event_type="FORMAT_CRITERIA_DRAFT_CREATED",
        actor_id=payload.lecturer_id,
        entity_type="format_criteria_set",
        entity_id=criteria_set.id,
        request_id=request_id,
        input_snapshot=payload.model_dump(mode="json"),
    )
    await session.commit()
    return await _criteria_response(session, criteria_set)


async def draft_from_documents(
    session: AsyncSession, payload: FormatDraftFromDocumentsCreate, request_id: str | None
) -> FormatCriteriaSetResponse:
    await _validate_course_lecturer(session, payload.course_id, payload.lecturer_id)
    documents = (
        await session.scalars(
            select(Document).where(
                Document.id.in_(payload.document_ids),
                Document.course_id == payload.course_id,
            )
        )
    ).all()
    if len(documents) != len(set(payload.document_ids)):
        raise AppError("FORMAT_SOURCE_DOCUMENT_NOT_FOUND", "Có tài liệu nguồn không thuộc học phần.", status_code=422)
    chunks = (
        await session.scalars(
            select(DocumentChunk)
            .where(DocumentChunk.document_id.in_(payload.document_ids))
            .order_by(DocumentChunk.document_id, DocumentChunk.chunk_index)
        )
    ).all()
    criteria: list[FormatCriterionInput] = []
    seen: set[tuple[str, str]] = set()
    for chunk in chunks:
        text = chunk.content
        for label, property_name, pattern in (
            ("Lề trái", "margins_cm", r"lề\s+trái\D{0,30}(\d+(?:[.,]\d+)?)\s*cm"),
            ("Lề phải", "margins_cm", r"lề\s+phải\D{0,30}(\d+(?:[.,]\d+)?)\s*cm"),
            ("Số trang", "page_count", r"(?:tối đa|không quá|từ)\s*(\d+)\s*trang"),
        ):
            match = re.search(pattern, text, flags=re.IGNORECASE)
            if not match:
                continue
            value = float(match.group(1).replace(",", "."))
            key = (property_name, label)
            if key in seen:
                continue
            seen.add(key)
            expected: dict[str, Any]
            if property_name == "margins_cm":
                expected = {"property": property_name, "value": {"left" if "trái" in label.lower() else "right": value}, "tolerance": 0.15}
            else:
                expected = {"property": property_name, "operator": "max", "value": int(value)}
            criteria.append(
                FormatCriterionInput(
                    name=label,
                    description=f"Giá trị được trích nháp từ tài liệu: {label.lower()} {value}.",
                    rule_type=FormatCriterionRuleType.MEASUREMENT,
                    expected_value=expected,
                    source_citation=text[:500],
                    source_document_id=chunk.document_id,
                )
            )
    if payload.include_template_structure:
        headings = [chunk.heading for chunk in chunks if chunk.heading]
        headings = list(dict.fromkeys(str(value) for value in headings))
        if headings:
            criteria.append(
                FormatCriterionInput(
                    name="Tuân thủ cấu trúc theo mẫu",
                    description="Khung đề mục được trích nháp từ cấu trúc tài liệu nguồn; giảng viên phải xác nhận trước khi dùng.",
                    rule_type=FormatCriterionRuleType.TEMPLATE,
                    expected_value={"headings": headings, "allow_extra": True, "strict_names": True},
                    source_citation="Khung đề mục trích từ tài liệu nguồn.",
                    source_document_id=documents[0].id,
                )
            )
    return await create_draft(
        session,
        FormatCriteriaDraftCreate(
            course_id=payload.course_id,
            lecturer_id=payload.lecturer_id,
            title=payload.title,
            criteria=criteria,
            source_document_ids=payload.document_ids,
        ),
        request_id,
    )


async def list_criteria_sets(session: AsyncSession, course_id: str) -> list[FormatCriteriaSetResponse]:
    items = (
        await session.scalars(
            select(FormatCriteriaSet)
            .where(FormatCriteriaSet.course_id == course_id)
            .order_by(FormatCriteriaSet.version.desc())
        )
    ).all()
    return [await _criteria_response(session, item) for item in items]


async def approve_criteria_set(
    session: AsyncSession, criteria_set_id: str, lecturer_id: str, request_id: str | None
) -> FormatCriteriaSetResponse:
    criteria_set = await session.get(FormatCriteriaSet, criteria_set_id)
    if criteria_set is None:
        raise AppError("FORMAT_CRITERIA_SET_NOT_FOUND", "Không tìm thấy bộ tiêu chí.", status_code=404)
    await _validate_course_lecturer(session, criteria_set.course_id, lecturer_id)
    if criteria_set.status != FormatCriteriaSetStatus.DRAFT:
        raise AppError("FORMAT_CRITERIA_NOT_DRAFT", "Chỉ bộ tiêu chí nháp mới có thể được duyệt.", status_code=409)
    previous = (
        await session.scalars(
            select(FormatCriteriaSet).where(
                FormatCriteriaSet.course_id == criteria_set.course_id,
                FormatCriteriaSet.status == FormatCriteriaSetStatus.ACTIVE,
            )
        )
    ).all()
    for item in previous:
        item.status = FormatCriteriaSetStatus.SUPERSEDED
    criteria_set.status = FormatCriteriaSetStatus.ACTIVE
    criteria_set.approved_by = lecturer_id
    criteria_set.approved_at = datetime.now(UTC)
    add_audit_event(
        session,
        event_type="FORMAT_CRITERIA_SET_APPROVED",
        actor_id=lecturer_id,
        entity_type="format_criteria_set",
        entity_id=criteria_set.id,
        request_id=request_id,
        output_snapshot={"version": criteria_set.version, "superseded_ids": [item.id for item in previous]},
    )
    await session.commit()
    return await _criteria_response(session, criteria_set)


async def get_active_criteria(session: AsyncSession, course_id: str) -> FormatCriteriaSet | None:
    return await session.scalar(
        select(FormatCriteriaSet).where(
            FormatCriteriaSet.course_id == course_id,
            FormatCriteriaSet.status == FormatCriteriaSetStatus.ACTIVE,
        )
    )


def _document_structure(filename: str, payload: bytes) -> dict[str, Any]:
    extension = Path(filename).suffix.lower()
    if extension == ".docx":
        document = DocxDocument(io.BytesIO(payload))
        headings = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip() and paragraph.style and paragraph.style.name.casefold().startswith("heading")
        ]
        sections = list(document.sections)
        margins = []
        for section in sections:
            margins.append({
                "top": round(section.top_margin.cm, 2), "bottom": round(section.bottom_margin.cm, 2),
                "left": round(section.left_margin.cm, 2), "right": round(section.right_margin.cm, 2),
            })
        font_names: list[str] = []
        font_sizes: list[float] = []
        line_spacings: list[float] = []
        for paragraph in document.paragraphs:
            spacing = paragraph.paragraph_format.line_spacing
            if isinstance(spacing, (int, float)):
                line_spacings.append(float(spacing))
            for run in paragraph.runs:
                if run.font.name:
                    font_names.append(run.font.name)
                if run.font.size:
                    font_sizes.append(round(run.font.size.pt, 2))
        tables = [[cell.text.strip() for cell in table.rows[0].cells] for table in document.tables if table.rows]
        return {
            "file_type": "DOCX", "page_count": None, "scan": False, "headings": headings,
            "margins_cm": margins, "font_names": list(Counter(font_names).keys()),
            "font_sizes_pt": sorted(set(font_sizes)), "line_spacings": sorted(set(line_spacings)),
            "table_headers": tables,
        }
    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(payload))
        page_texts = [(page.extract_text() or "").strip() for page in reader.pages]
        scan = not any(page_texts)
        lines = [line.strip() for text in page_texts for line in text.splitlines() if line.strip()]
        headings = [
            line for line in lines
            if re.match(r"^(?:\d+(?:\.\d+)*[.)]?\s+|[A-ZÀ-ỴĐ][A-ZÀ-ỴĐ\s]{5,}$)", line)
        ][:80]
        return {
            "file_type": "PDF", "page_count": len(reader.pages), "scan": scan, "headings": headings,
            "margins_cm": None, "font_names": None, "font_sizes_pt": None,
            "line_spacings": None, "table_headers": [],
        }
    raise AppError("FORMAT_FILE_TYPE_UNSUPPORTED", "Chỉ hỗ trợ file .docx hoặc .pdf.", status_code=422)


def _unmeasurable(criterion: FormatCriterion, reason: str) -> FormatCriterionResult:
    return FormatCriterionResult(
        criterion_id=criterion.id, name=criterion.name, status=FormatCriterionResultStatus.UNMEASURABLE,
        message=reason, expected_value=criterion.expected_value, source_citation=criterion.source_citation,
    )


def _measurement_result(criterion: FormatCriterion, structure: dict[str, Any]) -> FormatCriterionResult:
    expected = criterion.expected_value
    property_name = str(expected.get("property", ""))
    if property_name not in SUPPORTED_MEASUREMENTS:
        return _unmeasurable(criterion, "Tiêu chí này chưa nằm trong năng lực đo tự động hiện có.")
    if structure["file_type"] == "PDF" and property_name != "page_count":
        reason = "PDF là bản scan nên không có lớp văn bản đáng tin cậy." if structure["scan"] else "PDF không cung cấp đủ thông tin định dạng để đo tiêu chí này đáng tin cậy."
        return _unmeasurable(criterion, reason)
    structure_key = {
        "font_name": "font_names",
        "font_size_pt": "font_sizes_pt",
        "line_spacing": "line_spacings",
    }.get(property_name, property_name)
    actual: Any = structure.get(structure_key)
    if property_name == "page_count" and actual is None:
        return _unmeasurable(criterion, "Tệp DOCX không có thông tin số trang ổn định để đo mà không render tài liệu.")
    if actual in (None, [], {}):
        return _unmeasurable(criterion, "Không trích xuất được thuộc tính này một cách đáng tin cậy.")
    tolerance = float(expected.get("tolerance", 0.1))
    desired = expected.get("value")
    passed = False
    if property_name == "page_count":
        operator = expected.get("operator", "equals")
        passed = {"max": actual <= desired, "min": actual >= desired, "equals": actual == desired}.get(operator, False)
    elif property_name == "margins_cm":
        samples = actual
        if any(
            max(float(sample.get(side, 0)) for sample in samples)
            - min(float(sample.get(side, 0)) for sample in samples)
            > tolerance * 2
            for side in ("top", "bottom", "left", "right")
        ):
            return _unmeasurable(
                criterion,
                "Các phần đại diện có lề chênh lệch quá nhiều nên phép đo không đủ tin cậy.",
            )
        passed = bool(samples) and all(
            all(abs(float(sample.get(side, -999)) - float(value)) <= tolerance for side, value in desired.items())
            for sample in samples
        )
    elif property_name == "font_name":
        passed = any(_fold(str(value)) == _fold(str(desired)) for value in actual)
    elif property_name in {"font_size_pt", "line_spacing"}:
        passed = any(abs(float(value) - float(desired)) <= tolerance for value in actual)
    return FormatCriterionResult(
        criterion_id=criterion.id, name=criterion.name,
        status=FormatCriterionResultStatus.PASS if passed else FormatCriterionResultStatus.FAIL,
        message="Giá trị đo được khớp tiêu chí." if passed else "Giá trị đo được không khớp tiêu chí đã duyệt.",
        expected_value=expected, actual_value={property_name: actual}, source_citation=criterion.source_citation,
    )


def _presence_result(criterion: FormatCriterion, structure: dict[str, Any]) -> FormatCriterionResult:
    expected = criterion.expected_value
    kind = expected.get("kind")
    desired = expected.get("value")
    candidates = structure["headings"] if kind == "heading" else structure["table_headers"] if kind == "table_columns" else None
    if candidates is None:
        return _unmeasurable(criterion, "Tiêu chí hiện diện này chưa nằm trong năng lực đo tự động hiện có.")
    if structure["scan"]:
        return _unmeasurable(criterion, "PDF là bản scan nên không thể đọc cấu trúc một cách đáng tin cậy.")
    desired_items = [str(item) for item in desired] if isinstance(desired, list) else [str(desired)]
    normalized_expected = [_fold(item) for item in desired_items]
    normalized_candidates = [_fold(" ".join(item) if isinstance(item, list) else item) for item in candidates]
    if any(value in normalized_candidates for value in normalized_expected):
        status, message = FormatCriterionResultStatus.PASS, "Đã tìm thấy cấu trúc khớp rõ ràng."
    else:
        partial = any(
            set(value.split()) & set(candidate.split())
            for value in normalized_expected for candidate in normalized_candidates
        )
        status, message = (
            (FormatCriterionResultStatus.AMBIGUOUS, "Tìm thấy cấu trúc có vẻ liên quan nhưng không đủ rõ để kết luận.")
            if partial else (FormatCriterionResultStatus.FAIL, "Không tìm thấy cấu trúc bắt buộc.")
        )
    return FormatCriterionResult(
        criterion_id=criterion.id, name=criterion.name, status=status, message=message,
        expected_value=expected, actual_value={"candidates": candidates}, source_citation=criterion.source_citation,
    )


def _template_result(criterion: FormatCriterion, structure: dict[str, Any]) -> FormatCriterionResult:
    expected = criterion.expected_value
    template = [str(item) for item in expected.get("headings", [])]
    headings = [str(item) for item in structure["headings"]]
    if not template or structure["scan"] or not headings:
        return _unmeasurable(criterion, "Không trích xuất được khung đề mục đủ tin cậy để so với mẫu.")
    actual_folded = [_fold(item) for item in headings]
    expected_folded = [_fold(item) for item in template]
    exact_indexes = [actual_folded.index(item) if item in actual_folded else -1 for item in expected_folded]
    if any(index < 0 for index in exact_indexes):
        partial = any(set(item.split()) & set(candidate.split()) for item in expected_folded for candidate in actual_folded)
        status = FormatCriterionResultStatus.AMBIGUOUS if partial else FormatCriterionResultStatus.FAIL
        message = "Có đề mục gần giống mẫu nhưng tên chưa khớp rõ ràng." if partial else "Thiếu đề mục bắt buộc của mẫu."
    elif exact_indexes != sorted(exact_indexes):
        status, message = FormatCriterionResultStatus.FAIL, "Đề mục có đủ nhưng sai thứ tự rõ ràng so với mẫu."
    elif not expected.get("allow_extra", True) and len(headings) != len(template):
        status, message = FormatCriterionResultStatus.FAIL, "Bài nộp có đề mục ngoài khung mẫu không được phép."
    else:
        status, message = FormatCriterionResultStatus.PASS, "Khung đề mục khớp mẫu đã duyệt."
    return FormatCriterionResult(
        criterion_id=criterion.id, name=criterion.name, status=status, message=message,
        expected_value=expected, actual_value={"headings": headings}, source_citation=criterion.source_citation,
    )


async def _record_metric(session: AsyncSession, course_id: str, status: FormatCheckOverallStatus) -> None:
    today = date.today()
    metric = await session.scalar(select(FormatCheckMetric).where(FormatCheckMetric.course_id == course_id, FormatCheckMetric.metric_date == today))
    if metric is None:
        metric = FormatCheckMetric(
            course_id=course_id,
            metric_date=today,
            total_checks=0,
            passed_checks=0,
            needs_fix_checks=0,
            unable_to_check_checks=0,
        )
        session.add(metric)
    metric.total_checks += 1
    if status == FormatCheckOverallStatus.PASS:
        metric.passed_checks += 1
    elif status == FormatCheckOverallStatus.NEEDS_FIX:
        metric.needs_fix_checks += 1
    else:
        metric.unable_to_check_checks += 1


async def _purge_expired_reviews(session: AsyncSession) -> None:
    expired = (await session.scalars(select(FormatReviewCase).where(FormatReviewCase.expires_at < datetime.now(UTC)))).all()
    for item in expired:
        await session.delete(item)


async def check_format(
    session: AsyncSession, *, course_id: str, actor_id: str, filename: str, payload: bytes, request_id: str | None
) -> FormatCheckResponse:
    if not filename or not payload:
        raise AppError("FORMAT_FILE_EMPTY", "Chưa nhận được file cần kiểm tra.", status_code=422)
    if len(payload) > get_settings().format_check_max_upload_bytes:
        raise AppError("FORMAT_FILE_TOO_LARGE", "File vượt quá dung lượng cho phép.", status_code=413)
    actor = await session.get(Actor, actor_id)
    if actor is None or actor.role != ActorRole.STUDENT:
        raise AppError("FORMAT_CHECK_ACTOR_NOT_FOUND", "Không tìm thấy sinh viên.", status_code=404)
    if await session.get(Course, course_id) is None:
        raise AppError("COURSE_NOT_FOUND", "Không tìm thấy học phần.", status_code=404)
    active = await get_active_criteria(session, course_id)
    if active is None:
        return FormatCheckResponse(
            overall_status=FormatCheckOverallStatus.UNABLE_TO_CHECK,
            message="Chưa có bộ tiêu chí hình thức được giảng viên duyệt cho học phần này.",
        )
    criteria = (await session.scalars(select(FormatCriterion).where(FormatCriterion.criteria_set_id == active.id).order_by(FormatCriterion.position))).all()
    usable = [item for item in criteria if item.automatable and not item.review_only]
    if not usable:
        return FormatCheckResponse(
            overall_status=FormatCheckOverallStatus.UNABLE_TO_CHECK,
            message="Bộ tiêu chí đang áp dụng chưa có tiêu chí nào hệ thống có thể kiểm tra tự động.",
            criteria_set_id=active.id, criteria_version=active.version, criteria_approved_at=active.approved_at,
        )
    try:
        structure = _document_structure(filename, payload)
    except AppError:
        raise
    except Exception as error:
        raise AppError("FORMAT_FILE_INVALID", "Không thể đọc file; file có thể bị hỏng hoặc được bảo vệ.", status_code=422) from error
    results: list[FormatCriterionResult] = []
    for criterion in criteria:
        if criterion.review_only or not criterion.automatable:
            results.append(_unmeasurable(criterion, "Tiêu chí này được giảng viên giữ để tự xem, không dùng để kiểm tra tự động."))
        elif criterion.rule_type == FormatCriterionRuleType.MEASUREMENT:
            results.append(_measurement_result(criterion, structure))
        elif criterion.rule_type == FormatCriterionRuleType.PRESENCE:
            results.append(_presence_result(criterion, structure))
        else:
            results.append(_template_result(criterion, structure))
    # Criteria intentionally marked review-only are shown transparently, but do
    # not turn an otherwise complete automatic check into an indeterminate one.
    usable_ids = {item.id for item in usable}
    statuses = {item.status for item in results if item.criterion_id in usable_ids}
    if FormatCriterionResultStatus.FAIL in statuses:
        overall, message = FormatCheckOverallStatus.NEEDS_FIX, "Có tiêu chí hình thức chưa đạt; hãy xem chi tiết để tự sửa trước khi nộp."
    elif FormatCriterionResultStatus.AMBIGUOUS in statuses:
        overall, message = FormatCheckOverallStatus.UNABLE_TO_CHECK, "Có tiêu chí cần giảng viên phán đoán; hồ sơ đã được gửi để xem xét."
    elif FormatCriterionResultStatus.UNMEASURABLE in statuses:
        overall, message = FormatCheckOverallStatus.UNABLE_TO_CHECK, "Một số tiêu chí không đo được một cách đáng tin cậy."
    else:
        overall, message = FormatCheckOverallStatus.PASS, "Các tiêu chí hình thức tự động đều đạt."
    review_case_id: str | None = None
    ambiguous = [item.model_dump(mode="json") for item in results if item.status == FormatCriterionResultStatus.AMBIGUOUS]
    if ambiguous:
        review = FormatReviewCase(
            id=new_id("format_review"), course_id=course_id, criteria_set_id=active.id, actor_id=actor_id,
            assigned_reviewer_id="lecturer-01", filename=filename, extracted_snapshot=structure,
            ambiguous_results=ambiguous, status=FormatReviewStatus.UNDER_REVIEW,
            expires_at=datetime.now(UTC) + timedelta(days=get_settings().format_review_retention_days),
        )
        session.add(review)
        review_case_id = review.id
        add_audit_event(session, event_type="FORMAT_REVIEW_CASE_CREATED", actor_id=actor_id, entity_type="format_review_case", entity_id=review.id, request_id=request_id, output_snapshot={"criteria_set_id": active.id})
    await _record_metric(session, course_id, overall)
    add_audit_event(session, event_type="FORMAT_CHECK_COMPLETED", actor_id=actor_id, entity_type="format_criteria_set", entity_id=active.id, request_id=request_id, output_snapshot={"overall_status": overall, "result_count": len(results)})
    await _purge_expired_reviews(session)
    await session.commit()
    return FormatCheckResponse(
        overall_status=overall, message=message, criteria_set_id=active.id, criteria_version=active.version,
        criteria_approved_at=active.approved_at, results=results, review_case_id=review_case_id,
    )


def _review_response(item: FormatReviewCase) -> FormatReviewCaseResponse:
    return FormatReviewCaseResponse(
        id=item.id, course_id=item.course_id, criteria_set_id=item.criteria_set_id, actor_id=item.actor_id,
        filename=item.filename, extracted_snapshot=item.extracted_snapshot, ambiguous_results=item.ambiguous_results,
        status=item.status, decision_reason=item.decision_reason, decided_at=item.decided_at,
        expires_at=item.expires_at, created_at=item.created_at,
    )


async def list_review_cases(session: AsyncSession, course_id: str, lecturer_id: str) -> list[FormatReviewCaseResponse]:
    await _validate_course_lecturer(session, course_id, lecturer_id)
    await _purge_expired_reviews(session)
    await session.commit()
    items = (await session.scalars(select(FormatReviewCase).where(FormatReviewCase.course_id == course_id).order_by(FormatReviewCase.created_at.desc()))).all()
    return [_review_response(item) for item in items]


async def list_student_review_cases(
    session: AsyncSession, course_id: str, actor_id: str
) -> list[FormatReviewCaseResponse]:
    actor = await session.get(Actor, actor_id)
    if actor is None or actor.role != ActorRole.STUDENT:
        raise AppError("FORMAT_CHECK_ACTOR_NOT_FOUND", "Không tìm thấy sinh viên.", status_code=404)
    await _purge_expired_reviews(session)
    await session.commit()
    items = (
        await session.scalars(
            select(FormatReviewCase)
            .where(
                FormatReviewCase.course_id == course_id,
                FormatReviewCase.actor_id == actor_id,
            )
            .order_by(FormatReviewCase.created_at.desc())
        )
    ).all()
    return [_review_response(item) for item in items]


async def decide_review_case(session: AsyncSession, review_id: str, payload: FormatReviewDecisionCreate, request_id: str | None) -> FormatReviewCaseResponse:
    item = await session.get(FormatReviewCase, review_id)
    if item is None:
        raise AppError("FORMAT_REVIEW_NOT_FOUND", "Không tìm thấy hồ sơ kiểm tra hình thức.", status_code=404)
    await _validate_course_lecturer(session, item.course_id, payload.reviewer_id)
    if item.status != FormatReviewStatus.UNDER_REVIEW:
        raise AppError("FORMAT_REVIEW_ALREADY_DECIDED", "Hồ sơ này đã có quyết định.", status_code=409)
    item.status = payload.decision
    item.decision_reason = payload.reason
    item.decided_at = datetime.now(UTC)
    add_audit_event(session, event_type="FORMAT_REVIEW_DECIDED", actor_id=payload.reviewer_id, entity_type="format_review_case", entity_id=item.id, request_id=request_id, input_snapshot=payload.model_dump(mode="json"))
    await session.commit()
    return _review_response(item)


async def format_metrics(session: AsyncSession, course_id: str, lecturer_id: str) -> list[FormatMetricResponse]:
    await _validate_course_lecturer(session, course_id, lecturer_id)
    items = (await session.scalars(select(FormatCheckMetric).where(FormatCheckMetric.course_id == course_id).order_by(FormatCheckMetric.metric_date.desc()).limit(31))).all()
    return [FormatMetricResponse.model_validate(item, from_attributes=True) for item in items]
